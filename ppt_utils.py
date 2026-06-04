import os
import comtypes.client
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pdf_utils import pdf_to_images


# ─── Existing conversion functions ────────────────────────────────────────────

def ppt_to_pdf(input_path, output_path):
    input_path = os.path.abspath(input_path)
    output_path = os.path.abspath(output_path)
    powerpoint = comtypes.client.CreateObject("Powerpoint.Application")
    try:
        deck = powerpoint.Presentations.Open(input_path)
        deck.SaveAs(output_path, 32)  # 32=ppSaveAsPDF
        deck.Close()
    finally:
        powerpoint.Quit()


def ppt_to_images(input_path, output_dir):
    input_path = os.path.abspath(input_path)
    output_dir = os.path.abspath(output_dir)
    app = comtypes.client.CreateObject("Powerpoint.Application")
    try:
        prs = app.Presentations.Open(input_path, WithWindow=False)
        base = os.path.splitext(os.path.basename(input_path))[0]
        for i, slide in enumerate(prs.Slides):
            out_name = os.path.join(output_dir, f"{base}_slide_{i+1}.jpg")
            slide.Export(out_name, "JPG")
        prs.Close()
    finally:
        app.Quit()


# ─── New: Create / Read / Edit ────────────────────────────────────────────────

# Default layout indices in a standard blank presentation
LAYOUT_TITLE_SLIDE = 0     # "Title Slide" layout
LAYOUT_TITLE_CONTENT = 1   # "Title and Content" layout
LAYOUT_BLANK = 6           # "Blank" layout


def create_ppt(output_path, slides_data=None, theme_color=None):
    """
    Create a new PowerPoint presentation.

    Args:
        output_path: .pptx file path to write.
        slides_data: List of dicts:
            {
              'title': str,
              'content': str,          # Body text (newlines become bullet points)
              'layout': 'title' | 'content' | 'blank',
              'title_color': '#RRGGBB' (optional),
              'bg_color': '#RRGGBB'    (optional, solid background)
            }
    """
    prs = Presentation()
    prs.slide_width = Inches(13.33)   # Widescreen 16:9
    prs.slide_height = Inches(7.5)

    for slide_data in (slides_data or []):
        layout_key = slide_data.get('layout', 'content')
        if layout_key == 'title':
            layout = prs.slide_layouts[LAYOUT_TITLE_SLIDE]
        elif layout_key == 'blank':
            layout = prs.slide_layouts[LAYOUT_BLANK]
        else:
            layout = prs.slide_layouts[LAYOUT_TITLE_CONTENT]

        slide = prs.slides.add_slide(layout)

        # Optional solid background colour
        bg_hex = slide_data.get('bg_color', '')
        if bg_hex:
            try:
                bg_hex = bg_hex.lstrip('#')
                r, g, b = int(bg_hex[0:2], 16), int(bg_hex[2:4], 16), int(bg_hex[4:6], 16)
                fill = slide.background.fill
                fill.solid()
                fill.fore_color.rgb = RGBColor(r, g, b)
            except Exception:
                pass

        # Title
        title_text = slide_data.get('title', '')
        content_text = slide_data.get('content', '')

        if slide.shapes.title and title_text:
            tf = slide.shapes.title.text_frame
            tf.text = title_text
            title_hex = slide_data.get('title_color', '')
            if title_hex:
                try:
                    title_hex = title_hex.lstrip('#')
                    r, g, b = int(title_hex[0:2], 16), int(title_hex[2:4], 16), int(title_hex[4:6], 16)
                    for para in tf.paragraphs:
                        for run in para.runs:
                            run.font.color.rgb = RGBColor(r, g, b)
                except Exception:
                    pass

        # Content / body
        if content_text:
            # Find the body placeholder (index 1 usually)
            body_ph = None
            for ph in slide.placeholders:
                if ph.placeholder_format.idx == 1:
                    body_ph = ph
                    break

            if body_ph is not None:
                tf = body_ph.text_frame
                tf.word_wrap = True
                lines = content_text.split('\n')
                for i, line in enumerate(lines):
                    if i == 0:
                        tf.text = line
                    else:
                        p = tf.add_paragraph()
                        p.text = line
            else:
                # Fallback: add a text box
                txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(12.0), Inches(5.5))
                tf = txBox.text_frame
                tf.word_wrap = True
                tf.text = content_text

    prs.save(output_path)
    print(f"Presentation created: {output_path}")


def read_ppt_info(input_path):
    """
    Read a .pptx and return a list of slide summaries.

    Returns:
        List of dicts:
            { 'index': int (0-based), 'title': str, 'content': str }
    """
    prs = Presentation(input_path)
    result = []
    for idx, slide in enumerate(prs.slides):
        title = ''
        content_parts = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            if shape.shape_id == slide.shapes.title.shape_id if slide.shapes.title else False:
                title = shape.text_frame.text
            else:
                content_parts.append(shape.text_frame.text)
        # Also grab title from placeholder
        if slide.shapes.title:
            title = slide.shapes.title.text
        result.append({
            'index': idx,
            'title': title,
            'content': '\n'.join(content_parts).strip(),
        })
    return result


def add_slide(input_path, output_path, title, content, layout='content'):
    """Add a new slide at the end of an existing presentation."""
    prs = Presentation(input_path)
    layout_key = layout
    if layout_key == 'title':
        slide_layout = prs.slide_layouts[LAYOUT_TITLE_SLIDE]
    elif layout_key == 'blank':
        slide_layout = prs.slide_layouts[LAYOUT_BLANK]
    else:
        slide_layout = prs.slide_layouts[LAYOUT_TITLE_CONTENT]

    slide = prs.slides.add_slide(slide_layout)
    if slide.shapes.title and title:
        slide.shapes.title.text = title

    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 1:
            ph.text_frame.text = content
            break

    prs.save(output_path)
    print(f"Slide added to: {output_path}")


def edit_slide(input_path, output_path, slide_idx, title, content):
    """
    Edit title and content of a specific slide (0-based index).
    """
    prs = Presentation(input_path)
    if slide_idx >= len(prs.slides):
        raise IndexError(f"Slide index {slide_idx} out of range ({len(prs.slides)} slides)")

    slide = prs.slides[slide_idx]
    if slide.shapes.title:
        slide.shapes.title.text = title

    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 1:
            ph.text_frame.text = content
            break

    prs.save(output_path)
    print(f"Slide {slide_idx} updated in: {output_path}")


def delete_slide(input_path, output_path, slide_idx):
    """Delete a slide by 0-based index."""
    prs = Presentation(input_path)
    if slide_idx >= len(prs.slides):
        raise IndexError(f"Slide index {slide_idx} out of range")
    xml_slides = prs.slides._sldIdLst
    prs.slides._sldIdLst.remove(xml_slides[slide_idx])
    prs.save(output_path)
    print(f"Slide {slide_idx} deleted from: {output_path}")
